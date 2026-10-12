import pandas as pd, json
R='/home/user/FTO2_NC_handover/'
S='/tmp/claude-0/-home-user-FTO2-NC-handover/3e9a7317-eb66-5172-89e1-6e435294d804/scratchpad/'
m=pd.read_csv(S+'merged.csv',dtype={'notification_number':str},low_memory=False)
lt={c['notification_number']:c['long_text'] for i in range(1,35) for c in json.load(open(S+f'batches/b{i:02d}.json'))}
keep=['notification_number','seq','programme','ligne','perimetre','date_creation','code','fin','text','n_def_nc',
 'groupe','sous_groupe','rca_nom','defaut_cable','provenance',
 'llm_categorie','llm_sous_categorie','llm_sous_code','llm_rca','llm_mineur','llm_confiance','llm_symptome','llm_raison',
 'llm_defaut_cable','llm_provenance','llm_code_conflict','llm_deux_defauts','llm_contexte_utile']
d=m[keep].rename(columns={'groupe':'parser_categorie','sous_groupe':'parser_sous_categorie','rca_nom':'parser_rca_nom','defaut_cable':'parser_defaut_cable','provenance':'parser_provenance','code':'code_ecrit','text':'snippet'})
from importlib.util import spec_from_file_location, module_from_spec
sp=spec_from_file_location('p',R+'01_parser/nc_defect_parser.py');P=module_from_spec(sp);sp.loader.exec_module(P)
rca_nom=pd.read_csv(R+'04_reference/sous_categories.csv',encoding='utf-8-sig').set_index('sous_categorie').rca_nom
d.insert(d.columns.get_loc('llm_rca')+1,'llm_rca_nom',d.llm_sous_categorie.map(rca_nom).fillna('Non classé'))
d['accord_sous_cat']=d.parser_sous_categorie==d.llm_sous_categorie
d['accord_categorie']=d.parser_categorie==d.llm_categorie
d.to_csv(R+'05_outputs/fto2_defects_llm_longtext.csv',index=False,encoding='utf-8-sig')
f=d[d.perimetre=='FTO2']
dis=f[~f.accord_sous_cat].copy()
dis['long_text']=dis.notification_number.map(lt).str.slice(0,1500)
dis=dis.sort_values(['n_def_nc','notification_number','seq'])
dis=dis[['notification_number','seq','n_def_nc','code_ecrit','snippet','long_text','parser_categorie','parser_sous_categorie','llm_categorie','llm_sous_categorie','llm_symptome','llm_raison','llm_confiance']]
dis['verdict (P=parser, L=LLM, A=autre)']='';dis['commentaire']=''
cc=f[f.llm_code_conflict][['notification_number','seq','code_ecrit','snippet','llm_sous_categorie','llm_raison']]
summ=pd.DataFrame({'Parser v4.8':f.parser_categorie.value_counts(),'LLM texte long':f.llm_categorie.value_counts()}).fillna(0).astype(int)
summ['Écart']=summ['LLM texte long']-summ['Parser v4.8']
summ=summ.loc[[c for c in P.GROUPES if c in summ.index]]
sub=pd.DataFrame({'Parser v4.8':f.parser_sous_categorie.value_counts(),'LLM texte long':f.llm_sous_categorie.value_counts()}).fillna(0).astype(int)
sub['Écart']=sub['LLM texte long']-sub['Parser v4.8']; sub=sub.sort_values('Écart')
readme=pd.DataFrame({'Read me':[
 'Re-labelling of every parsed defect (3,098; 2,768 in FTO2 perimeter) by Claude Sonnet, reading the defect snippet AND the full NC long text (+ title, SAP codes, station, author dept).',
 'Done inside the Claude session (no external API), 12 Oct 2026. Same 9 categories / sub-categories / RCA as parser v4.8. Rules: written code wins; root cause first, on the exact element at fault; E1-E4; provenance only from explicit wording.',
 'Root cause rule: a defect caused by lack of securing is labelled as the securing problem (e.g. gaine pincée dans le collier -> Frettage / colliers – pose); the observed consequence is in llm_symptome.',
 'Manual review after the run: 5 collar cases aligned to Frettage; 30 "câble manquant, à fournir" moved back from Amont to Manquant – câble (E4 needs explicit never-supplied wording).',
 'Sheet Désaccords: the 303 FTO2 defects where LLM and parser differ (single-defect NCs first). Fill the verdict column to measure which is right.',
 'Sheet Conflits code: 19 defects where the written code contradicts the text (code kept, real cause in llm_raison).',
 'Full table: 05_outputs/fto2_defects_llm_longtext.csv. Method + scripts: 07_scripts/llm_longtext/.']})
with pd.ExcelWriter(R+'05_outputs/FTO2_llm_longtext_review.xlsx') as w:
    readme.to_excel(w,sheet_name='Read me',index=False)
    summ.to_excel(w,sheet_name='Catégories')
    sub.to_excel(w,sheet_name='Sous-catégories')
    dis.to_excel(w,sheet_name='Désaccords',index=False)
    cc.to_excel(w,sheet_name='Conflits code',index=False)
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
wb=load_workbook(R+'05_outputs/FTO2_llm_longtext_review.xlsx')
for ws in wb:
    ws.freeze_panes='B2' if ws.title!='Read me' else None
    for c in ws[1]: c.font=Font(bold=True)
    for col in ws.columns:
        L=max(len(str(c.value or '')) for c in list(col)[:200]); ws.column_dimensions[col[0].column_letter].width=min(max(10,L*0.9),60)
ws=wb['Désaccords']
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
ws.auto_filter.ref=ws.dimensions
wb['Read me'].column_dimensions['A'].width=160
wb.save(R+'05_outputs/FTO2_llm_longtext_review.xlsx')
print(len(d),len(dis),len(cc)); print(summ)
