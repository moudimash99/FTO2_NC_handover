"""Parts to price for the CNQ part value: one row per part, with how often it comes back.

The cost model counts the value of the part on rejections (REJ), rejections with a rework order (REJ_RW), exchanges (EXC)
and scraps (SCR) where FTO2 is liable, read from a `part_value` column that does not exist yet (README_handover gap 1).
The SAP article on the NC is mostly the assembly being built (a VU, a routing), so the part is read from the NC text:
its reference when one is written, else its FIN, else the part as the NC names it ("To identify").

Reads 05_outputs/fto2_dash_nc.csv (decision, liability, phase...) and the raw exports in 06_input_data (title, long
text, SAP article, quantity). Writes 05_outputs/CNQ_parts_to_price.xlsx: Read me, Parts to price (price to fill in),
NC detail (the per-NC part value the model needs), Not charged to FTO2.

    python3 -I 07_scripts/cnq/parts_to_price.py
"""
import os
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT = os.path.join(ROOT, "05_outputs", "CNQ_parts_to_price.xlsx")
EXPORTS = ("fto2_dashboard_data_12.csv", "fto2_dashboard_data_missing_ids_3.csv")
PART_RULES = ["REJ", "REJ_RW", "EXC", "SCR"]

# part references as written in the NC texts, in the order they appear
PN = [
    r"\b(?:LOT)?([FVMLD]\d{13}(?:S\d\d)?)\b",                                   # Airbus article, 14 characters
    r"\b((?:NSA|NSB|ABS|ASNA|ASNE|NAS|DAN)\d[0-9A-Z]*(?:-[0-9A-Z]+)*|(?:EN|MS|AN)\d{4,}[0-9A-Z]*(?:-[0-9A-Z]+)*)\b",  # standard parts
    r"\b(E0\d{2}[0-9A-Z]*(?:-\d+)?)\b",                                          # equipment (E0...)
    r"\b((?!10|33|0000)\d{10})\b",                                               # 10-digit references (not orders, phones, zeros)
]
FIN = re.compile(r"\b(\d{1,4}[A-Z]{2,3}\d{0,3}(?:-[A-D])?)\b")
# SAP articles that are the rejected part itself rather than the assembly or routing being worked on
PART_ARTICLE = re.compile(r"\bLABEL\b|WIRE BUNDLE|^VTP\b", re.I)


def body(t):
    """Long text without the author stamps, phone numbers and SAP boilerplate."""
    t = t if isinstance(t, str) else ""
    t = re.sub(r"\*?\s*\d\d\.\d\d\.\d{4} \d\d:\d\d:\d\d \w+ \([^)]*\)", " ", t)
    t = re.sub(r"T[ée]l\.\s*[\d#~ ]+", " ", t)
    t = re.sub(r"Non ?Conformity \d+ : Type of nonconformity : [^*]*\*", " ", t)
    t = re.sub(r"Task Completed.*$", " ", t)
    return re.sub(r"(\s*--)+", " ", re.sub(r"\s+", " ", t)).strip()


def refs_in(text):
    up = text.upper()
    up = re.sub(r"\b(ABS|NSA|ASNA|NAS) (\d)", r"\1\2", up)                    # "ABS 1424-3-10"
    up = re.sub(r"\b([FV]\d{3})\.(\d{5})\.(\d{5})\b", r"\1\2\3", up)              # "V924.19487.00000"
    found = []
    for pat in PN:
        found += [(m.start(), m.group(1)) for m in re.finditer(pat, up)]
    seen, out = set(), []
    for _, r in sorted(found):
        if r not in seen:
            seen.add(r)
            out.append(r)
    return out


def fins_in(text):
    out = []
    for f in FIN.findall(text.upper()):
        if not re.search(r"(VU|VB)$", f) and f not in out:
            out.append(f)
    return out


def load():
    raw = pd.concat([pd.read_csv(os.path.join(ROOT, "06_input_data", f), low_memory=False, dtype=str) for f in EXPORTS])
    raw["k"] = raw.notification_number.str.lstrip("0")
    raw = raw.drop_duplicates("k")
    nc = pd.read_csv(os.path.join(ROOT, "05_outputs", "fto2_dash_nc.csv"), dtype=str, keep_default_na=False, encoding="utf-8-sig")
    nc = nc[nc.rule.isin(PART_RULES)].copy()
    nc["k"] = nc.nc.str.lstrip("0")
    m = nc.merge(raw, on="k", how="left", suffixes=("", "_raw"))
    rows = []
    for r in m.itertuples(index=False):
        short = (r.description_short if isinstance(r.description_short, str) else "").strip()
        text = short + " | " + body(r.description_long_text)
        refs, fins = refs_in(text), fins_in(text)
        q = pd.to_numeric(r.defect_quantity, errors="coerce")
        art_desc = r.article_version_description if isinstance(r.article_version_description, str) else ""
        if isinstance(r.article_version, str) and PART_ARTICLE.search(art_desc):
            key, kind = r.article_version, "SAP article (the part itself)"
        elif refs:
            key, kind = refs[0], "Reference in text"
        elif fins:
            key, kind = "FIN %s" % fins[0], "FIN only"
        else:
            key, kind = re.sub(r"\s+", " ", re.sub(r"MSN ?\d+\.?", "", short)).strip(" .").upper() or "(no title)", "To identify"
        rows.append({
            "nc": r.nc, "decision": r.rule, "decision_label": r.decision_label, "liable": r.on_aec == "1", "liability": r.liability,
            "phase": r.phase, "line": r.line, "product": r.product, "msn": r.msn, "created": r.created[:10],
            "qty": float(q) if q == q and q > 0 else 1.0, "key": key, "kind": kind, "part_named": short,
            "refs": " ".join(refs), "fins": " ".join(fins),
            "sap_article": r.article_version if isinstance(r.article_version, str) else "",
            "sap_article_desc": r.article_version_description if isinstance(r.article_version_description, str) else "",
            "text": text[:600],
        })
    return pd.DataFrame(rows)


def groups(t):
    order = {"SAP article (the part itself)": 0, "Reference in text": 0, "FIN only": 1, "To identify": 2}
    out = []
    for key, g in t.groupby("key"):
        dec = g.decision.value_counts()
        out.append({
            "key": key, "kind": g.kind.iloc[0], "part": g.part_named.mode().iloc[0], "ncs": len(g), "msns": g.msn.nunique(),
            "qty": g.qty.sum(), "decisions": ", ".join("%s %d" % (d, n) for d, n in dec.items()),
            "products": ", ".join(sorted(set(g["product"]))), "lines": ", ".join(sorted(set(g.line) - {""})),
            "first": g.created.min(), "last": g.created.max(), "nc_list": " ".join(g.sort_values("created").nc),
        })
    return sorted(out, key=lambda r: (-r["ncs"], order[r["kind"]], r["key"]))


# ------------------------------------------------------------------ workbook
ARIAL = "Arial"
F = Font(name=ARIAL, size=10)
FB = Font(name=ARIAL, size=10, bold=True)
FH = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
FIN_INPUT = Font(name=ARIAL, size=10, color="0000FF")
FILL_H = PatternFill("solid", fgColor="2F3B52")
FILL_IN = PatternFill("solid", fgColor="FFFF00")
FILL_TOT = PatternFill("solid", fgColor="E8EBF0")
THIN = Border(bottom=Side(style="thin", color="C8CDD6"))
EUR = '#,##0.00;(#,##0.00);"-"'
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


def table(ws, headers, widths, rows, wrap_cols=()):
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font, cell.fill, cell.alignment = FH, FILL_H, Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = w
    for i, vals in enumerate(rows, 2):
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font, cell.border = F, THIN
            cell.alignment = WRAP if c in wrap_cols else TOP
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 30


def build(t):
    a = t[t.liable].copy()
    g = groups(a)
    wb = Workbook()

    # --- Read me
    rm = wb.active
    rm.title = "Read me"
    n_ref = sum(1 for r in g if r["kind"] == "Reference in text")
    n_fin = sum(1 for r in g if r["kind"] == "FIN only")
    n_id = sum(1 for r in g if r["kind"] == "To identify")
    n_art = sum(1 for r in g if r["kind"] == "SAP article (the part itself)")
    lines = [
        ("CNQ – parts to price", FB),
        ("", F),
        ("What this is", FB),
        ("The cost model counts the value of the part on rejections (REJ), rejections with a rework order (REJ_RW), exchanges (EXC) "
         "and scraps (SCR) where FTO2 is liable. Today that value is 0 because no price is known.", F),
        ("%d NCs need a part price. They come down to %d parts: %d whose SAP article is the part itself, %d with a reference "
         "written in the NC, %d with only a FIN, %d to identify from the NC's wording." % (len(a), len(g), n_art, n_ref, n_fin, n_id), F),
        ("Source: NC exports of 12 Oct 2026 (fto2_dashboard_data_12 and _missing_ids_3), decisions and liability as in the dashboard "
         "(NCs created in 2026, every phase).", F),
        ("", F),
        ("How the part was found", FB),
        ("The reference is read from the NC title and long text (NSA, ABS, ASNA, NAS, EN, E0…, 14-character Airbus articles, "
         "10-digit references). When none is written, the FIN is used; otherwise the part is listed as the NC names it.", F),
        ("The SAP article on the NC is mostly the assembly or the routing being worked on (ELEC ASSY 295VU, INTEGRATION 740VU, "
         "Gamme domestique 743VU INCOMING, Intégration harnais 1101VB sur MIP), not the rejected part: the same articles carry "
         "rework NCs, where nothing is thrown away. When the article is the part itself (a functional label, a harness), it is "
         "used as the part. Otherwise it is shown in 'NC detail' for reference only.", F),
        ("This is an automatic reading of the texts: check the reference before asking for a price.", F),
        ("", F),
        ("What to fill in", FB),
        ("Sheet 'Parts to price', yellow cells only: Unit price (€) and Price source. Part value = unit price × quantity.", F),
        ("Sheet 'NC detail' picks the price up by itself. Its 'Part value (€)' column is what the cost model reads as part_value.", F),
        ("Example of the expected format (not a real price): Unit price 65.40 · Price source 'SAP standard price, Oct 2026'.", F),
        ("", F),
        ("To decide before using the values", FB),
        ("Quantity is the SAP defect quantity (1 on almost every NC). Check it when several parts were replaced.", F),
        ("An exchanged part returned under warranty may not cost its full price.", F),
        ("REJ_RW NCs are also charged the rework time: keep both only if the part was really thrown away and reworked.", F),
        ("NCs where FTO2 is not liable cost 0 in the model; they are listed in 'Not charged to FTO2' for completeness (%d NCs)."
         % int((~t.liable).sum()), F),
    ]
    for i, (txt, font) in enumerate(lines, 1):
        c = rm.cell(row=i, column=1, value=txt)
        c.font, c.alignment = font, Alignment(wrap_text=True, vertical="top")
    rm["A1"].font = Font(name=ARIAL, size=14, bold=True)
    rm.column_dimensions["A"].width = 120

    # --- Parts to price
    ws = wb.create_sheet("Parts to price")
    headers = ["Part reference / FIN / name", "Found as", "Part (as named in the NC)", "NCs", "MSNs", "Quantity", "Decisions",
               "Products", "Lines", "First NC", "Last NC", "NC numbers", "Unit price (€)", "Price source", "Part value (€)"]
    widths = [24, 16, 40, 7, 7, 9, 18, 18, 10, 11, 11, 30, 13, 26, 14]
    table(ws, headers, widths, [[r["key"], r["kind"], r["part"], r["ncs"], r["msns"], r["qty"], r["decisions"], r["products"],
                                  r["lines"], r["first"], r["last"], r["nc_list"], None, None, None] for r in g], wrap_cols=(3, 12))
    last = len(g) + 1
    for i in range(2, last + 1):
        for col in ("M", "N"):
            ws[f"{col}{i}"].fill, ws[f"{col}{i}"].font = FILL_IN, FIN_INPUT
        ws[f"M{i}"].number_format = EUR
        ws[f"O{i}"] = f'=IF(M{i}="","",M{i}*F{i})'
        ws[f"O{i}"].number_format = EUR
        ws[f"F{i}"].number_format = "0"
    tot = last + 2
    ws[f"A{tot}"] = "Total"
    ws[f"D{tot}"] = f"=SUM(D2:D{last})"
    ws[f"F{tot}"] = f"=SUM(F2:F{last})"
    ws[f"M{tot}"] = f'=COUNT(M2:M{last})&" of "&COUNTA(A2:A{last})&" priced"'
    ws[f"O{tot}"] = f"=SUM(O2:O{last})"
    for col in "ABCDEFGHIJKLMNO":
        ws[f"{col}{tot}"].font, ws[f"{col}{tot}"].fill = FB, FILL_TOT
    ws[f"F{tot}"].number_format = "0"
    ws[f"O{tot}"].number_format = EUR
    ws.auto_filter.ref = f"A1:O{last}"
    ws["M1"].comment = Comment("Fill in: price of one part in euros (SAP standard price or purchase price).", "CNQ")
    ws["F1"].comment = Comment("SAP defect quantity of each NC, summed. 1 on almost every NC: check when several parts were replaced.", "CNQ")

    # --- NC detail
    wd = wb.create_sheet("NC detail")
    headers = ["NC", "Decision", "Decision label", "Phase", "Line", "Product", "MSN", "Created", "Quantity", "Part reference / FIN / name",
               "Part (as named in the NC)", "References in text", "FIN in text", "SAP article of the NC", "SAP article description",
               "Unit price (€)", "Part value (€)", "NC text (extract)"]
    widths = [11, 9, 22, 8, 8, 11, 7, 11, 9, 24, 36, 26, 12, 20, 28, 13, 14, 80]
    d = a.sort_values(["key", "created"])
    table(wd, headers, widths, [[r.nc, r.decision, r.decision_label, r.phase, r.line, r.product, r.msn, r.created, r.qty, r.key,
                                  r.part_named, r.refs, r.fins, r.sap_article, r.sap_article_desc, None, None, r.text]
                                 for r in d.itertuples()], wrap_cols=(11, 18))
    lastd = len(d) + 1
    look = "INDEX('Parts to price'!$M$2:$M${0},MATCH(J{1},'Parts to price'!$A$2:$A${0},0))"
    for i in range(2, lastd + 1):
        wd[f"P{i}"] = f'=IFERROR(IF({look.format(last, i)}="","",{look.format(last, i)}),"")'
        wd[f"Q{i}"] = f'=IF(P{i}="","",P{i}*I{i})'
        wd[f"P{i}"].number_format = wd[f"Q{i}"].number_format = EUR
        wd[f"I{i}"].number_format = "0"
    wd[f"A{lastd + 2}"] = "Total"
    wd[f"Q{lastd + 2}"] = f"=SUM(Q2:Q{lastd})"
    wd[f"Q{lastd + 2}"].number_format = EUR
    for col in "ABCDEFGHIJKLMNOPQR":
        wd[f"{col}{lastd + 2}"].font, wd[f"{col}{lastd + 2}"].fill = FB, FILL_TOT
    wd.auto_filter.ref = f"A1:R{lastd}"

    # --- Not charged to FTO2
    wn = wb.create_sheet("Not charged to FTO2")
    o = t[~t.liable].sort_values(["key", "created"])
    table(wn, ["NC", "Decision", "Liability", "Phase", "Product", "Created", "Part reference / FIN / name", "Part (as named in the NC)"],
          [11, 9, 30, 8, 11, 11, 24, 40],
          [[r.nc, r.decision, r.liability, r.phase, r.product, r.created, r.key, r.part_named] for r in o.itertuples()], wrap_cols=(3, 8))
    wn.auto_filter.ref = f"A1:H{len(o) + 1}"

    wb.active = 1
    wb.save(OUT)
    return a, g


def main():
    t = load()
    a, g = build(t)
    kinds = pd.Series([r["kind"] for r in g]).value_counts().to_dict()
    print("NCs with a part value:", len(t), "| FTO2 liable:", len(a), "| parts:", len(g), kinds,
          "| parts on 2+ NCs:", sum(1 for r in g if r["ncs"] >= 2))
    print("written", os.path.relpath(OUT, ROOT))


if __name__ == "__main__":
    main()
