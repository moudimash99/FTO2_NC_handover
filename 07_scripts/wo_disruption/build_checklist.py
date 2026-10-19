# -*- coding: utf-8 -*-
"""
Build the CNQ investigation workbook (checklist + one tab per item) from the analysis pickles.

    python3 -I build_checklist.py <P> <repo> <disrup.csv> <out.xlsx>

<P> is the folder written by load.py and the other scripts of this folder (see README.md).
"""
import re, sys
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

D, REPO, DISRUP, OUT = sys.argv[1:5]
wo = pd.read_pickle(D + "/wo.pkl"); wo["h"] = pd.to_numeric(wo.labor_standard_value, errors="coerce").fillna(0) / 100
fam = pd.read_pickle(D + "/s03_family.pkl")
nr = pd.read_pickle(D + "/nc_rw.pkl"); n26, n25 = nr["2026"], nr["2025"]
X = pd.read_pickle(D + "/sheet_data.pkl"); R, c26 = X["R"], X["c26"]
ds = pd.read_pickle(D + "/ds2.pkl"); ds["tl"] = pd.to_numeric(ds.time_lost)
raw = pd.read_csv(DISRUP, dtype=str, encoding="utf-8-sig", keep_default_na=False, na_values=[""])

# ------------------------------------------------------------------ styles
F = "Arial"
f_base, f_bold, f_head = Font(name=F, size=10), Font(name=F, size=10, bold=True), Font(name=F, size=10, bold=True, color="FFFFFF")
f_title, f_sub = Font(name=F, size=14, bold=True, color="1F3864"), Font(name=F, size=10, italic=True, color="595959")
f_input_val = Font(name=F, size=10, color="0000FF")
fill_head, fill_input = PatternFill("solid", fgColor="1F3864"), PatternFill("solid", fgColor="FFF2CC")
fill_hi, fill_med, fill_low = PatternFill("solid", fgColor="F8CBAD"), PatternFill("solid", fgColor="FFE699"), PatternFill("solid", fgColor="E2EFDA")
fill_total = PatternFill("solid", fgColor="D9E1F2")
thin = Side(style="thin", color="BFBFBF"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
clean = lambda v: ILLEGAL_CHARACTERS_RE.sub(" ", v) if isinstance(v, str) else v
one_line = lambda s, n=400: "" if not isinstance(s, str) else " ".join(s.split())[:n]
wb = Workbook()

def sheet(name, title, sub):
    ws = wb.create_sheet(name)
    ws["A1"] = title; ws["A1"].font = f_title
    ws["A2"] = sub; ws["A2"].font = f_sub
    ws.sheet_view.showGridLines = False
    return ws

def para(ws, row, text, bold=False, ncols=10):
    ws.cell(row=row, column=1, value=clean(text)).font = f_bold if bold else f_base
    ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    ws.row_dimensions[row].height = max(15, 15 * (1 + len(text) // 150))
    return row + 1

def table(ws, row, df, widths, inputs=(), yn=(), status=(), fmt=None, filt=True):
    """Write df at row; columns named in `inputs` are yellow cells to fill. Returns (first data row, last data row)."""
    fmt = fmt or {}
    cols = list(df.columns)
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=j, value=c); cell.font = f_head; cell.fill = fill_head; cell.alignment = Alignment(wrap_text=True, vertical="center"); cell.border = box
        ws.column_dimensions[get_column_letter(j)].width = widths[j - 1]
    ws.row_dimensions[row].height = 30
    r0 = row + 1
    for i, rec in enumerate(df.itertuples(index=False), r0):
        for j, v in enumerate(rec, 1):
            if isinstance(v, float) and np.isnan(v): v = None
            if isinstance(v, (np.integer,)): v = int(v)
            if isinstance(v, (np.floating,)): v = float(v)
            cell = ws.cell(row=i, column=j, value=clean(v)); cell.font = f_base; cell.alignment = wrap; cell.border = box
            if cols[j - 1] in fmt: cell.number_format = fmt[cols[j - 1]]
            if cols[j - 1] in inputs: cell.fill = fill_input; cell.font = f_input_val
    r1 = r0 + max(len(df), 1) - 1
    if len(df) == 0:
        for j, c in enumerate(cols, 1):
            if c in inputs: ws.cell(row=r0, column=j).fill = fill_input
    for c in list(yn) + list(status):
        j = cols.index(c) + 1; L = get_column_letter(j)
        dv = DataValidation(type="list", formula1='"Y,N,?"' if c in yn else '"To do,In progress,Done"', allow_blank=True)
        ws.add_data_validation(dv); dv.add(f"{L}{r0}:{L}{r1}")
    if filt: ws.auto_filter.ref = f"A{row}:{get_column_letter(len(cols))}{r1}"
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    return r0, r1

def total_row(ws, row, label, formulas, ncols):
    ws.cell(row=row, column=1, value=label).font = f_bold
    for j in range(1, ncols + 1):
        ws.cell(row=row, column=j).fill = fill_total; ws.cell(row=row, column=j).border = box
    for col, f, nf in formulas:
        c = ws.cell(row=row, column=col, value=f); c.font = f_bold
        if nf: c.number_format = nf

def col(df, name): return get_column_letter(list(df.columns).index(name) + 1)

# ================================================================== Figures (first, referenced by others)
wsF = sheet("Figures", "CNQ figures as of 19 October 2026", "Model cnq-2026.10.19 (alert only when an operator raises one; connection errors 35 then 15 min; internal NCs x 2). Blue = values from the analysis, black = formulas.")
wsF["A4"] = "Hourly rate used for hours (EUR/h)"; wsF["A4"].font = f_bold
wsF["C4"] = 50; wsF["C4"].font = f_input_val; wsF["C4"].fill = fill_input; wsF["C4"].number_format = "0.00"
wsF["D4"] = "Official France rate used by the CNQ model (RATE_STANDARD). Change it here and every EUR figure in this workbook follows."; wsF["D4"].font = f_sub
RATE = "Figures!$C$4"

# ================================================================== D block (option B)
wsD = sheet("D_Block", "Disruption block (option B): quality aléas that did not lead to an NC",
            "MES disruptions of type 'QLS - Probleme qualite', Jan-Sep 2026, with no NC on the same work order within -2 h..+72 h. Aléas that led to an NC are inside the NC cost (alert step), not here.")
qb = R["blockB"].sort_values("ot").copy()
dfD = pd.DataFrame({
    "Disruption": qb.disruption_id, "Opened": qb.ot.dt.tz_convert("Europe/Paris").dt.strftime("%Y-%m-%d %H:%M"),
    "Month (yyyymm)": qb.ot.dt.tz_convert("Europe/Paris").dt.strftime("%Y%m").astype(int), "Line": qb.line, "Station": qb.physical_station,
    "Work order": qb.work_order, "Time lost (min)": qb.tl.astype(float), "What the operator wrote": qb.comment.map(one_line),
    "Answer": qb.solution.fillna("").str.replace("solve$$", "", regex=False).map(one_line),
    "Possible ERMA NC (MIP / 1101VB)": np.where(qb.wf201_like, "Y", ""),
    "Keep in block? (Y/N)": "Y", "Comment": ""})
r = para(wsD, 4, "Pre-filled Y for every row. Put N for an aléa that should not count (for example, its NC exists but is written by ERMA, which is out of scope). The totals below and in Figures follow.", ncols=12)
d0, d1 = table(wsD, 6, dfD, [11, 15, 9, 13, 14, 12, 9, 60, 40, 12, 10, 25], inputs=["Keep in block? (Y/N)", "Comment"], yn=["Keep in block? (Y/N)"])
cK, cT, cM = col(dfD, "Keep in block? (Y/N)"), col(dfD, "Time lost (min)"), col(dfD, "Month (yyyymm)")
total_row(wsD, d1 + 1, "Kept in block", [(list(dfD.columns).index("Time lost (min)") + 1, f'=SUMIFS({cT}{d0}:{cT}{d1},{cK}{d0}:{cK}{d1},"<>N")', "#,##0"),
                                          (2, f'=COUNTIF({cK}{d0}:{cK}{d1},"<>N")&" aléas"', None)], len(dfD.columns))
D_RANGE = dict(T=f"D_Block!${cT}${d0}:${cT}${d1}", K=f"D_Block!${cK}${d0}:${cK}${d1}", M=f"D_Block!${cM}${d0}:${cM}${d1}")

# ================================================================== M1 quality-like aléas of other kinds
wsM1 = sheet("M1_QualityLike", "M1. Aléas of other kinds whose text describes a quality problem (no NC followed)",
             "Logistics, work-order and line disruptions whose comment reads like a quality problem. Grouped so you can answer yes / no per group instead of reading 107 rows.")
g3 = R["m1"].sort_values(["category", "ot"]).copy()
cats = ["Cable damaged, cut too short or duplicated", "Workmanship error on the line (crimping, forgotten, wrong fitting)",
        "Part or component broken, damaged or faulty", "Wrong or non-conforming part or document supplied", "Other wording"]
ex = {c: " | ".join(g3[g3.category == c].comment.map(lambda s: one_line(s, 110)).head(3)) for c in cats}
r = para(wsM1, 4, "Answer Y or N in the yellow column for each group. Rows of a group marked Y are added to the disruption block (see Figures). The full list is below the table.", ncols=8)
hdr = ["Group", "Aléas", "Hours lost", "Three examples", "Count as quality? (Y/N)", "Comment"]
for j, h in enumerate(hdr, 1):
    c = wsM1.cell(row=6, column=j, value=h); c.font = f_head; c.fill = fill_head; c.border = box; c.alignment = Alignment(wrap_text=True)
for j, w in enumerate([45, 9, 10, 90, 14, 30], 1): wsM1.column_dimensions[get_column_letter(j)].width = w
L0 = 15  # list header row
dfM1 = pd.DataFrame({"Disruption": g3.disruption_id, "Opened": g3.ot.dt.tz_convert("Europe/Paris").dt.strftime("%Y-%m-%d"),
                     "Month (yyyymm)": g3.ot.dt.tz_convert("Europe/Paris").dt.strftime("%Y%m").astype(int), "Kind": g3.type_desc, "Line": g3.line,
                     "Time lost (min)": g3.tl.astype(float), "Group": g3.category, "What the operator wrote": g3.comment.map(one_line), "Counted": ""})
m0, m1 = L0 + 1, L0 + len(dfM1)
for i, cname in enumerate(cats, 7):
    wsM1.cell(row=i, column=1, value=cname)
    wsM1.cell(row=i, column=2, value=f'=COUNTIF($G${m0}:$G${m1},A{i})')
    wsM1.cell(row=i, column=3, value=f'=SUMIFS($F${m0}:$F${m1},$G${m0}:$G${m1},A{i})/60').number_format = "0.0"
    wsM1.cell(row=i, column=4, value=clean(ex[cname]))
    for j in range(1, 7):
        c = wsM1.cell(row=i, column=j); c.border = box; c.alignment = wrap; c.font = f_base
    wsM1.cell(row=i, column=5).fill = fill_input; wsM1.cell(row=i, column=6).fill = fill_input
dv = DataValidation(type="list", formula1='"Y,N,?"', allow_blank=True); wsM1.add_data_validation(dv); dv.add("E7:E11")
total_row(wsM1, 12, "Counted as quality (groups marked Y)", [(2, f'=COUNTIF($I${m0}:$I${m1},"Y")', None), (3, f'=SUMIFS($F${m0}:$F${m1},$I${m0}:$I${m1},"Y")/60', "0.0")], 6)
table(wsM1, L0, dfM1, [11, 11, 9, 38, 14, 9, 45, 90, 9], filt=True)
wsM1.column_dimensions["A"].width = 45
for i in range(m0, m1 + 1):
    wsM1.cell(row=i, column=9, value=f'=IF(IFERROR(INDEX($E$7:$E$11,MATCH(G{i},$A$7:$A$11,0)),"")="Y","Y","N")').font = f_base
wsM1.freeze_panes = "A7"
M1_RANGE = dict(T=f"M1_QualityLike!$F${m0}:$F${m1}", C=f"M1_QualityLike!$I${m0}:$I${m1}", M=f"M1_QualityLike!$C${m0}:$C${m1}")

# ================================================================== H3 customer returns
wsH3 = sheet("H3_CustomerReturns", "H3. Customer returns: the actual cost, including the work done at the customer",
             "Customer-return NCs where FTO2 is liable, created in 2026. We are the direct cause, so every cost we can find counts, even outside the plant.")
cr = R["cr"].copy().join(n26.set_index("notification_number")[["creation_time", "msn", "description_short", "workflow_list"]], how="left")
cr = cr.sort_values("creation_time")
dfH3 = pd.DataFrame({
    "NC": cr.index, "Created": cr.creation_time.astype(str).str[:10], "Raised at plant": cr.plant, "Programme": cr.programme, "MSN": cr.msn.astype(str).str.replace(".0", "", regex=False),
    "What was found": cr.description_short.map(one_line), "Decision": cr.cnq_rule, "Defect lines": cr.cnq_defect_lines,
    "Model cost now (EUR)": cr.cnq_cost_total.round(2), "of which rework at FTO2 times (EUR)": cr.cnq_cost_work.round(2),
    "Rework order(s) linked": cr.rework_orders, "Planned h on that order (Sept file)": cr.planned_h.round(2),
    "Repair hours at the customer": None, "Analysis / 8D hours": None, "Transport, kit, logistics (EUR)": None, "Penalties, other (EUR)": None,
    "Source of the figures": "", "Cost with findings (EUR)": None})
r = para(wsH3, 4, "What the model counts today: writing the NC, technical support, re-check, and the rework estimated at FTO2's times (no operator alert since 19 Oct). "
         "What it does not count: the repair time actually spent at the customer, the analysis (8D), transport, kits, penalties. Fill what you find in the yellow columns; "
         "'Cost with findings' replaces the model's rework by the real repair hours when you enter them, and adds the rest.", ncols=12)
r = para(wsH3, 5, "What we already know: FNZ1 returns (cockpit 'mauvais backshell') get an FTO2 rework order that only checks the kit sent to Montoir-de-Bretagne (0 h planned). "
         "Two FFF1 returns ('312VU endommagé') are linked to REC rework orders at the FAL (0.5 and 0.8 h planned). In September the quality team's A350 order booked 15 h of "
         "'Gestion NC retours clients + analyse' for 8 A350 returns.", ncols=12)
h0, h1 = table(wsH3, 7, dfH3, [11, 10, 8, 8, 7, 40, 8, 7, 11, 11, 22, 10, 11, 10, 11, 11, 25, 12],
               inputs=["Repair hours at the customer", "Analysis / 8D hours", "Transport, kit, logistics (EUR)", "Penalties, other (EUR)", "Source of the figures"],
               fmt={"Model cost now (EUR)": "#,##0.00", "of which rework at FTO2 times (EUR)": "#,##0.00", "Cost with findings (EUR)": "#,##0.00", "Planned h on that order (Sept file)": "0.0"})
cI, cJ, cM_, cN, cO, cP, cR = [col(dfH3, n) for n in ["Model cost now (EUR)", "of which rework at FTO2 times (EUR)", "Repair hours at the customer", "Analysis / 8D hours",
                                                      "Transport, kit, logistics (EUR)", "Penalties, other (EUR)", "Cost with findings (EUR)"]]
for i in range(h0, h1 + 1):
    c = wsH3[f"{cR}{i}"]
    c.value = f"={cI}{i}-IF(ISNUMBER({cM_}{i}),{cJ}{i}-{cM_}{i}*{RATE},0)+N({cN}{i})*{RATE}+N({cO}{i})+N({cP}{i})"
    c.number_format = "#,##0.00"; c.font = f_base
total_row(wsH3, h1 + 1, "Total", [(list(dfH3.columns).index(n) + 1, f"=SUM({col(dfH3, n)}{h0}:{col(dfH3, n)}{h1})", "#,##0.00")
                                   for n in ["Model cost now (EUR)", "of which rework at FTO2 times (EUR)", "Cost with findings (EUR)"]], len(dfH3.columns))
H3_TOT = (f"H3_CustomerReturns!${cI}${h1 + 1}", f"H3_CustomerReturns!${cR}${h1 + 1}", len(dfH3))

# ================================================================== H1 rework without NC
wsH1 = sheet("H1_ReworkNoNC", "H1. Rework orders on the FTO2 shop floor with no NC behind them (September 2026)",
             "S03 rework orders on FTO2 work centres (REC, flight test, regularisation and A400M excluded) that no NC links to and that name no NC number.")
def kind(t):
    t = t.lower()
    if re.search(r"gestion nc|alerte qualit|ad-hoc|accompagnement", t): return "Quality team time (see H2)"
    if re.search(r"fagot|débiter du", t): return "Cable cutting for stock (see H4)"
    if re.search(r"formboard", t): return "Formboard check after a definition change"
    if re.search(r"nc cable|suite à|suite a", t): return "Cable remake that refers to an NC without its number"
    if re.search(r"impri|câble|cable|débit|plaquette|étiquette", t): return "Cable or label remake, no NC reference"
    if re.search(r"test|continuit|emi|sonn|remise en conformit", t): return "Rework after test, no NC written"
    return "Other rework"
h1o = fam[(fam.family == "FTO2 shop floor") & ~fam.traced].copy()
h1o["cat"] = h1o.t.map(kind)
h1o = h1o.sort_values(["cat", "h"], ascending=[True, False])
dfH1 = pd.DataFrame({"Rework order": h1o.index, "Programme": h1o.prog, "MSN": h1o.msn, "Work centre(s)": h1o.wcd.map(lambda s: one_line(s, 120)),
                     "Planned h": h1o.h.round(2), "Group (my reading)": h1o.cat, "Named in disruption": h1o.in_disruption.fillna(""),
                     "What the order says": h1o.t.map(lambda s: one_line(s, 350)),
                     "Non-quality? (Y/N)": "", "Should an NC have been written? (Y/N)": "", "Comment": ""})
r = para(wsH1, 4, "For each order (or each group, using the filter): is it non-quality, and should an NC exist? Rework after test and cable remakes with no NC are the most likely 'missing NCs'. "
         "Planned hours are SAP standard values in IM (1/100 h), converted to hours.", ncols=11)
a0, a1 = table(wsH1, 6, dfH1, [12, 9, 7, 30, 8, 30, 12, 80, 11, 13, 25], inputs=["Non-quality? (Y/N)", "Should an NC have been written? (Y/N)", "Comment"],
               yn=["Non-quality? (Y/N)", "Should an NC have been written? (Y/N)"], fmt={"Planned h": "0.0"})
cH = col(dfH1, "Planned h"); cY = col(dfH1, "Non-quality? (Y/N)")
total_row(wsH1, a1 + 1, "Total planned h", [(5, f"=SUM({cH}{a0}:{cH}{a1})", "0.0"), (6, f'="Marked non-quality: "&TEXT(SUMIFS({cH}{a0}:{cH}{a1},{cY}{a0}:{cY}{a1},"Y"),"0.0")&" h"', None)], len(dfH1.columns))
H1_TOT = (f"H1_ReworkNoNC!${cH}${a1 + 1}", f'SUMIFS(H1_ReworkNoNC!${cH}${a0}:${cH}${a1},H1_ReworkNoNC!${cY}${a0}:${cY}${a1},"Y")', len(dfH1))

# ================================================================== H2 quality team time
wsH2 = sheet("H2_QualityTime", "H2. Quality team time booked on rework orders (September 2026)",
             "Two S03 orders on quality work centres look like the quality team's time sheet for the month: one for the A350 line, one for the A330 line.")
q2 = wo[wo.order_number.isin(["1031243823", "1031225552"])].sort_values(["order_number", "operation_number"])
dfH2 = pd.DataFrame({"Order": q2.order_number, "Programme / MSN": q2.order_program + " / " + q2.order_msn.astype(str), "Work centre": q2.work_center,
                     "Op": q2.operation_number, "Operation": q2.operation_short_text, "Long text": q2.operation_long_text.map(lambda s: one_line(s, 120)),
                     "Planned h": q2.h.round(2), "NC-related? (Y/N)": "", "Already in the model's steps? (writing / support / none)": "", "Comment": ""})
r = para(wsH2, 4, "Questions: Is there one such order per line every month? Who books on it? Which operations are NC handling (and so CNQ), and which are normal quality work "
         "(QA qualification, production support)? Does any of it overlap the model's 'writing the NC' (30/60 min) or 'technical support' (60 min) steps?", ncols=10)
b0, b1 = table(wsH2, 6, dfH2, [12, 14, 11, 6, 40, 40, 9, 11, 22, 25], inputs=["NC-related? (Y/N)", "Already in the model's steps? (writing / support / none)", "Comment"],
               yn=["NC-related? (Y/N)"], fmt={"Planned h": "0.0"})
cH2, cY2 = col(dfH2, "Planned h"), col(dfH2, "NC-related? (Y/N)")
total_row(wsH2, b1 + 1, "Total planned h", [(7, f"=SUM({cH2}{b0}:{cH2}{b1})", "0.0"), (8, f'=SUMIFS({cH2}{b0}:{cH2}{b1},{cY2}{b0}:{cY2}{b1},"Y")', "0.0")], len(dfH2.columns))
wsH2.cell(row=b1 + 2, column=8, value="h marked NC-related").font = f_sub

# ================================================================== H4 cable cutting
wsH4 = sheet("H4_CableCutting", "H4. BF24 / BF20 cables cut for stock on rework orders (A330 print room, September 2026)",
             "37 S03 orders on work centre 4FEM0E14 asking for bundles of 50 cables. The PTVU line had raised disruptions about BF24 quantities in Elec ABC and missing BF24 cables.")
t = (wo.operation_short_text.fillna("") + " " + wo.operation_long_text.fillna("")).str.replace(r"\s+", " ", regex=True)
cut = wo[t.str.contains(r"fagot|débiter du|Débiter câbles BF|Débiter câble BF", case=False) & (wo.order_type == "S03")].copy(); cut["txt"] = t[cut.index]
oc = cut.groupby("order_number").agg(h=("h", "sum"), msn=("order_msn", "first"), day=("actual_start_execution_date", lambda x: str(min(x))[:10]), txt=("txt", " | ".join)).sort_values("day")
dfH4 = pd.DataFrame({"Order": oc.index, "Confirmed": oc.day, "MSN": oc.msn, "Planned h": oc.h.round(2), "Request": oc.txt.map(lambda s: one_line(s, 200)), "Comment": ""})
r = para(wsH4, 4, "Question: is this a consequence of non-quality (cables remade or replaced because of defects) or a definition / kitting problem (cut lists not giving the line the right cables)? "
         "Who to ask: methods (Elec ABC quantities), the A330 print room, the PTVU line leader. The related disruptions and NCs are listed further down.", ncols=6)
e0, e1 = table(wsH4, 6, dfH4, [12, 11, 7, 9, 110, 30], inputs=["Comment"], fmt={"Planned h": "0.0"}, filt=False)
total_row(wsH4, e1 + 1, "Total planned h", [(4, f"=SUM(D{e0}:D{e1})", "0.0")], 6)
rr = e1 + 3
wsH4.cell(row=rr, column=1, value="Related disruptions (mention BF24 / BF20)").font = f_bold
bf = ds[ds.comment.fillna("").str.contains(r"BF ?2[04]", case=False)].sort_values("ot")
dfb = pd.DataFrame({"Disruption": bf.disruption_id, "Opened": bf.ot.dt.tz_convert("Europe/Paris").dt.strftime("%Y-%m-%d"), "Line": bf.line, "Time lost (min)": bf.tl.astype(float),
                    "What the operator wrote": bf.comment.map(lambda s: one_line(s, 300)), "Kind": bf.type_desc})
for j, c in enumerate(dfb.columns, 1):
    x = wsH4.cell(row=rr + 1, column=j, value=c); x.font = f_head; x.fill = fill_head; x.border = box
for i, rec in enumerate(dfb.itertuples(index=False), rr + 2):
    for j, v in enumerate(rec, 1):
        x = wsH4.cell(row=i, column=j, value=clean(v)); x.font = f_base; x.alignment = wrap; x.border = box
rr = rr + 3 + len(dfb)
wsH4.cell(row=rr, column=1, value="Related NCs (2026 NCs mentioning BF24 / BF20)").font = f_bold
bn = n26[n26.description_long_text.fillna("").str.contains(r"BF ?2[04]")].sort_values("creation_time")
dfn = pd.DataFrame({"NC": bn.notification_number, "Created": bn.creation_time.astype(str).str[:10], "Programme": bn.programme, "Type": bn.Type_NC_auto,
                    "What was found": bn.description_short.map(one_line), "Long text (extract)": bn.description_long_text.map(lambda s: one_line(s[max(0, s.find("BF") - 120):], 250))})
for j, c in enumerate(dfn.columns, 1):
    x = wsH4.cell(row=rr + 1, column=j, value=c); x.font = f_head; x.fill = fill_head; x.border = box
for i, rec in enumerate(dfn.itertuples(index=False), rr + 2):
    for j, v in enumerate(rec, 1):
        x = wsH4.cell(row=i, column=j, value=clean(v)); x.font = f_base; x.alignment = wrap; x.border = box

# ================================================================== H5 NCs reworked with no rework order
wsH5 = sheet("H5_NCsNoOrder", "H5. Reworked NCs with no rework order",
             "Decision 'rework', FTO2 liable, raised at FTO2, but no rework order in the NC's linked_document (an AW_/WO_/RS_ number different from its own work order).")
rows = []
for yr, n in [("2026", n26), ("2025", n25)]:
    z = n[(n.cnq_rule == "RWK") & (n.cnq_on_aec == 1) & (n.plant == "FTO2") & ~n.has_rw & (pd.to_datetime(n.creation_time, utc=True, errors="coerce").dt.year == int(yr))]
    rows.append(pd.DataFrame({"NC": z.notification_number.astype(str), "Year": yr, "Created": z.creation_time.astype(str).str[:10], "NC type": z.Type_NC_auto,
                              "Workflow": z.workflow_list, "Station": z.station, "Written by": z.user_department, "What was found": z.description_short.map(one_line),
                              "linked_document": z.linked_document, "Own work order": z.work_order.astype(str).str.replace(".0", "", regex=False)}))
dfH5 = pd.concat(rows, ignore_index=True)
dfH5["Fixed on the spot? (Y/N)"] = ""; dfH5["Order exists but link missing? (Y/N)"] = ""; dfH5["Comment"] = ""
r = para(wsH5, 4, "96-97% of reworked NCs get their own rework order. For these few, check in SAP whether the rework was done on the spot (no order needed) or whether an order exists "
         "but was never linked to the NC. If orders exist unlinked, the 'rework order' rule misses them.", ncols=13)
table(wsH5, 6, dfH5, [11, 6, 10, 13, 13, 9, 26, 35, 20, 12, 11, 13, 25], inputs=["Fixed on the spot? (Y/N)", "Order exists but link missing? (Y/N)", "Comment"],
      yn=["Fixed on the spot? (Y/N)", "Order exists but link missing? (Y/N)"])

# ================================================================== H6 rework orders with no planned time
wsH6 = sheet("H6_OrdersNoTime", "H6. Rework orders with no planned time (September 2026)",
             "S03 rework orders whose operations carry no standard value. If planned time is the truth, these orders cost 0, which cannot be right.")
z = R["noplan"].copy()
z["nc"] = z.nc_link_2026.fillna("")
z = z.join(n26.set_index("notification_number")[["Type_NC_auto", "description_short"]], on="nc")
z = z.sort_values(["family", "traced"], ascending=[True, False])
dfH6 = pd.DataFrame({"Rework order": z.index, "Family": z.family, "Programme": z.prog, "MSN": z.msn, "Linked NC": z.nc, "NC type": z.Type_NC_auto.fillna(""),
                     "NC: what was found": z.description_short.map(one_line), "Work centre(s)": z.wcd.map(lambda s: one_line(s, 100)),
                     "Operations": z.t.map(lambda s: one_line(s, 250)), "Why no time? (your finding)": "", "Comment": ""})
r = para(wsH6, 4, "Questions: why do these orders have no standard value (operations created without time, closed without confirmation, time booked elsewhere)? "
         "Until known, the CNQ uses the model's estimate for NCs whose order has no planned time. 63 of the 249 shop-floor rework orders (25%) are in this case.", ncols=11)
table(wsH6, 6, dfH6, [12, 16, 9, 7, 11, 13, 30, 30, 70, 30, 25], inputs=["Why no time? (your finding)", "Comment"])

# ================================================================== H7 disruption export
wsH7 = sheet("H7_DisruptExport", "H7. Is the disruption export complete? (2,500 rows)",
             "The file has exactly 2,500 rows, sorted newest first: the pattern of an export cut at a row limit.")
o = pd.to_datetime(raw.open_date_time)
facts = [("Rows in the file", len(raw)), ("First disruption", str(o.min())), ("Last disruption", str(o.max())),
         ("File order", "newest first (30 Sep 23:31 on the first row)"), ("NCs created 29 Dec 2025 - 2 Jan 2026 (plant closed?)", "0"),
         ("First NC of 2026", "5 Jan 2026 (same day as the first disruption)"), ("Usual time of the first disruption of a day", "about 07:43"),
         ("Time of the first disruption on 5 Jan", "08:57 (9 disruptions that day; a usual Monday has about 14)")]
wsH7.column_dimensions["A"].width = 48; wsH7.column_dimensions["B"].width = 60; wsH7.column_dimensions["C"].width = 30
for i, (k, v) in enumerate(facts, 4):
    wsH7.cell(row=i, column=1, value=k).font = f_bold; wsH7.cell(row=i, column=2, value=v).font = f_base
r = 4 + len(facts) + 1
for s in ["How to check (any one is enough):",
          "1. Re-export the disruptions with only 1 Jan - 31 Mar 2026 selected. If 5 January gets more than 9 rows, or rows before 08:57, the first export was cut.",
          "2. Look in the export tool (Skywise / Contour) for a row limit or a 'first N rows' setting.",
          "3. Count the disruptions in MES for 5 January 2026 and compare with 9."]:
    wsH7.cell(row=r, column=1, value=s).font = f_bold if s.endswith(":") else f_base; wsH7.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3); r += 1
r += 1
for j, h in enumerate(["Check", "Result (your finding)", "Export complete? (Y/N)"], 1):
    c = wsH7.cell(row=r, column=j, value=h); c.font = f_head; c.fill = fill_head; c.border = box
for k, nm in enumerate(["Re-export Jan-Mar: rows on 5 Jan", "Row limit in the export tool", "MES count for 5 Jan"], r + 1):
    wsH7.cell(row=k, column=1, value=nm).font = f_base
    for j in (2, 3): wsH7.cell(row=k, column=j).fill = fill_input; wsH7.cell(row=k, column=j).border = box
dv = DataValidation(type="list", formula1='"Y,N,?"', allow_blank=True); wsH7.add_data_validation(dv); dv.add(f"C{r + 1}:C{r + 3}")
r += 5
wsH7.cell(row=r, column=1, value="Disruptions per day, first two weeks of the file").font = f_bold
dd = o.dt.date.value_counts().sort_index().head(10)
for i, (k, v) in enumerate(dd.items(), r + 1):
    wsH7.cell(row=i, column=1, value=str(k)).font = f_base; wsH7.cell(row=i, column=2, value=int(v)).font = f_base
wsH7.cell(row=r + 12, column=1, value="Time columns: time_lost is typed by the operator (minutes). time_to_get_ack / time_to_get_fix do not match the timestamps (1-2 h time-zone shift; time_to_get_fix is 0 for 64%).").font = f_sub

# ================================================================== M2 rework orders naming NCs we don't have
wsM2 = sheet("M2_NCsNotHeld", "M2. Rework orders that name an NC we do not have (September 2026)",
             "Their text says 'Suite à la NC 2183…' but that NC is in neither the 2026 dashboard export nor the 2025 file: probably outside the dashboard scope.")
m2 = fam[fam.status == "E. names an NC we DON'T have"].copy().sort_values("h", ascending=False)
dfM2 = pd.DataFrame({"Rework order": m2.index, "Family": m2.family, "Programme": m2.prog, "MSN": m2.msn, "NC named": m2.nc_in_text.map(", ".join), "Planned h": m2.h.round(2),
                     "What the order says": m2.t.map(lambda s: one_line(s, 250)), "Plant of the NC": "", "Liability": "", "Written by (ERMA?)": "", "Comment": ""})
r = para(wsM2, 4, "Look up each NC in Skywise: plant, liability and user_department. If they are ERMA NCs or NCs of other plants where FTO2 is not liable, they are out of scope as decided. If FTO2 is liable, our NC export misses them.", ncols=11)
f0, f1 = table(wsM2, 6, dfM2, [12, 16, 9, 7, 12, 8, 70, 12, 12, 14, 25], inputs=["Plant of the NC", "Liability", "Written by (ERMA?)", "Comment"], fmt={"Planned h": "0.0"})
total_row(wsM2, f1 + 1, "Total planned h", [(6, f"=SUM(F{f0}:F{f1})", "0.0")], len(dfM2.columns))

# ================================================================== M3 exceptional NCs
wsM3 = sheet("M3_Exceptional", "M3. Exceptional NCs (to check yourself; not for the presentation)",
             "Reworked NCs whose rework order plans at least 3 h and more than 2.5 times the model's rework time (September 2026 orders).")
ex = R["exc"].join(n26.set_index("notification_number")[["description_short", "description_long_text"]])
dfM3 = pd.DataFrame({"NC": ex.index, "NC type": ex.Type_NC_auto, "Defect lines (parser)": ex.cnq_defect_lines, "Planned rework (h)": (ex.total / 60).round(1),
                     "Model rework (h)": (ex.model_new / 60).round(1), "What was found": ex.description_short.map(one_line),
                     "Long text (extract)": ex.description_long_text.map(lambda s: one_line(s[150:] if isinstance(s, str) else s, 300)), "Finding": ""})
r = para(wsM3, 4, "NC 218233692 ('Câbles non lasérisés') is one defect line for the model but its order plans 34 h of reprinting cables for six routes. The three others have no defect "
         "line found by the parser, so the model uses a flat 1 hour. Check whether the planned time is right and whether such NCs should be costed from their order.", ncols=8)
table(wsM3, 6, dfM3, [11, 12, 9, 10, 10, 35, 80, 30], inputs=["Finding"], fmt={"Planned rework (h)": "0.0", "Model rework (h)": "0.0"})

# ================================================================== Later
wsL = sheet("Later", "Kept for later, and points to keep in mind", "Not for tomorrow's checks, but not to forget.")
later = [
 ("Change 5 (kept for later)", "Two defect classes produced by parser v4.8 are missing from the CNQ time table and get the 10-min default: 'Mesure hors tolérance' (~25 lines a year) and 'Dommage – tresse de masse' (2-3 lines). Proposal: give them the time of 'Point de masse / mesure' (5 min) and 'Dommage – câble' (30 min). Effect: almost 0 EUR."),
 ("Point to keep in mind: alert on customer returns", "By the 'direct cause' rule, the customer's own operator alert should count too. Change 3 removed the alert from customer returns; putting it back adds 25 EUR per return (69 returns in 2026, about 1,700 EUR). Decide with H3."),
 ("Point to keep in mind: flat 30-min alert", "Aléas that led to an NC declare on average 41 min lost (median 30). The NC cost uses a flat 30 min. When the disruption link is in the pipeline, the declared time can replace it."),
 ("Point to keep in mind: 4 final-inspection NCs with an aléa", "4 final-inspection NCs did start with a quality aléa (4.6 h) but get no alert step since change 3, and are not in the disruption block either. Small; fixed when the disruption link is in the pipeline."),
 ("Point to keep in mind: responders' time", "Disruption 'time lost' is the operator's time only. The time of the people who answer the aléa (quality, technical support, logistics) is not counted anywhere."),
 ("Point to keep in mind: periods", "The disruption block covers Jan-Sep 2026 only (no 2025 disruptions in the data). The NC cost for 2026 is quoted for Feb-Sep because January is cut in the NC export."),
 ("Low: MSN 700/701 work (IEVB work centres)", "74 rework orders, 217 h in September on MSN 700/701 (e.g. 'Repose FTI sur bouteille oxy'). Confirm it is flight-test installation and out of scope."),
 ("After the presentation", "Full year of rework orders (re-check 15 min per connection error and x2 on internal NCs; use planned time per NC in the pipeline). Quality-team time sheets for every month. Repair cost of customer returns from TOF / FNZ1 / FFF1. Euros already booked by Finance on rework orders. Disruption link (work order + 72 h) in the pipeline."),
]
for j, (h, w) in enumerate([("Item", 40), ("Detail", 130)], 1):
    c = wsL.cell(row=4, column=j, value=h); c.font = f_head; c.fill = fill_head; wsL.column_dimensions[get_column_letter(j)].width = w
for i, (k, v) in enumerate(later, 5):
    wsL.cell(row=i, column=1, value=k).font = f_bold; c = wsL.cell(row=i, column=2, value=v); c.font = f_base; c.alignment = wrap
    wsL.cell(row=i, column=1).alignment = wrap; wsL.row_dimensions[i].height = 45

# ================================================================== Figures (body)
def put(ws, row, label, val, fmt=None, blue=False, note=None, bold=False):
    ws.cell(row=row, column=1, value=label).font = f_bold if bold else f_base
    c = ws.cell(row=row, column=3, value=val); c.font = f_input_val if blue else (f_bold if bold else f_base)
    if fmt: c.number_format = fmt
    if note: ws.cell(row=row, column=4, value=note).font = f_sub
wsF.column_dimensions["A"].width = 62; wsF.column_dimensions["B"].width = 4; wsF.column_dimensions["C"].width = 16; wsF.column_dimensions["D"].width = 95
r = 6; wsF.cell(row=r, column=1, value="1. NC cost of non-quality (level 3)").font = f_bold; r += 1
for lab, v, note in [("Pack reference data, 5 Jan - 6 Jul 2026 (595 NCs)", 121503, "cnq_handoever/tests/expected_totals.json (parser v4.7 data)"),
                     ("2026 Feb-Sep, raised at FTO2 (695 NCs)", 142151, "Parser v4.8 on the dashboard export + missing IDs; January left out (export cut)"),
                     ("2025 Jan-Dec, raised at FTO2, without ERMA (818 NCs)", 168633, "2025 NC file; ERMA NCs left out to match the 2026 pipeline"),
                     ("Customer returns 2026, Jan-Sep (69 NCs), model today", 10250, "Raised at other plants, FTO2 liable; see H3 for the actual cost")]:
    put(wsF, r, lab, v, "#,##0", True, note); r += 1
put(wsF, r, "Customer returns 2026 with the findings entered in H3", f"={H3_TOT[1]}", "#,##0", note="Follows H3 'Cost with findings'"); r += 2
wsF.cell(row=r, column=1, value="2. Disruption block (option B): quality aléas with no NC").font = f_bold; r += 1
put(wsF, r, "Declared quality aléas kept (D_Block), Jan-Sep 2026: hours lost", f"=SUMIFS({D_RANGE['T']},{D_RANGE['K']},\"<>N\")/60", "#,##0.0"); rB1 = r; r += 1
put(wsF, r, "  of which Feb-Sep (same period as the 2026 NC cost)", f"=SUMIFS({D_RANGE['T']},{D_RANGE['K']},\"<>N\",{D_RANGE['M']},\">=202602\")/60", "#,##0.0"); rB2 = r; r += 1
put(wsF, r, "Other-kind aléas counted as quality (M1 groups marked Y): hours lost", f"=SUMIFS({M1_RANGE['T']},{M1_RANGE['C']},\"Y\")/60", "#,##0.0"); rB3 = r; r += 1
put(wsF, r, "Disruption block, Jan-Sep 2026 (EUR)", f"=(C{rB1}+C{rB3})*{RATE}", "#,##0", bold=True, note="Hours lost x rate. Shown beside the NC cost, never added inside it."); r += 2
wsF.cell(row=r, column=1, value="3. Known, not counted (scope or under investigation)").font = f_bold; r += 1
put(wsF, r, "REC rework orders at the FAL, September 2026: planned hours", 746, "#,##0", True, "Out of scope as decided (REC)"); rR = r; r += 1
put(wsF, r, "  in EUR", f"=C{rR}*{RATE}", "#,##0"); r += 1
put(wsF, r, "ERMA NCs 2025, FTO2 liable (184 NCs): level 3 EUR", 44472, "#,##0", True, "Out of scope as decided (ERMA); new model"); r += 1
put(wsF, r, "Shop-floor rework with no NC, September 2026: planned hours (H1)", f"={H1_TOT[0]}", "#,##0.0", note="Under investigation (H1)"); rH = r; r += 1
put(wsF, r, "  of which marked non-quality in H1: hours", f"={H1_TOT[1]}", "#,##0.0"); r += 1
put(wsF, r, "  of which marked non-quality in H1: EUR", f"=C{r - 1}*{RATE}", "#,##0"); r += 1

# ================================================================== Checklist
wsC = sheet("Checklist", "CNQ investigation checklist", "Built 19 October 2026 from the work orders (Sept 2026), the MES disruptions (Jan-Sep 2026) and the NC files (2025, 2026). One line per check; each has its own tab.")
items = [
 ("H1", "High", "Rework done without an NC (shop floor)", f"{H1_TOT[2]} rework orders, about 181 planned h in September, no NC behind them: quality-team time, cable cutting, rework after test with no NC, cable / label remakes with no NC reference, formboard checks.",
  "Which are non-quality? Which should have had an NC?", "Rework order texts in SAP; line leaders; quality", "H1_ReworkNoNC"),
 ("H2", "High", "Quality team time booked on rework orders", "2 orders, 70.5 h in September (A350 and A330 lines), incl. 15 h 'Gestion NC retours clients + analyse' and 4 h 'Alerte qualité'.",
  "Monthly? Which hours are NC handling? Overlap with the model's writing / support steps?", "Quality line managers (4QEH5I07, 4QEH1I07); SAP orders", "H2_QualityTime"),
 ("H3", "High", "Customer returns: actual cost incl. work at the customer", f"{H3_TOT[2]} returns in 2026, model cost about 10,250 EUR. Repair happens at the customer (FNZ1 kits, FFF1 via REC orders); analysis, transport, penalties not counted.",
  "Repair hours at the customer, analysis hours, transport / kits, penalties. Put the customer's alert back?", "Quality (customer returns), TOF / FNZ1 / FFF1 contacts, REC", "H3_CustomerReturns"),
 ("H4", "High", "BF24 / BF20 cable cutting for stock", "37 orders, 78 h, 2,600 cables cut in September for the PTVU line, after disruptions on Elec ABC quantities and missing BF24.",
  "Non-quality, or a definition / kitting problem?", "Methods (Elec ABC), A330 print room, PTVU line leader", "H4_CableCutting"),
 ("H5", "High", "Reworked NCs with no rework order", f"{len(dfH5)} NCs (2026 and 2025) decided 'rework' with nothing in linked_document.",
  "Fixed on the spot, or an order that was never linked?", "SAP: NC and orders on the same work order / MSN", "H5_NCsNoOrder"),
 ("H6", "High", "Rework orders with no planned time", f"{len(dfH6)} of 614 rework orders in September have no standard value (63 of 249 on the shop floor).",
  "Why no time? How should those NCs be costed?", "Methods / technical support who create the orders", "H6_OrdersNoTime"),
 ("H7", "High", "Disruption export complete?", "Exactly 2,500 rows, newest first, 5 Jan 08:57 to 30 Sep 23:31. Starts on the plant's restart day, so probably complete.",
  "Re-export Jan-Mar and compare 5 January; look for a row limit.", "Skywise / Contour export; MES", "H7_DisruptExport"),
 ("M1", "Medium", "Other aléas that look like quality", f"{len(dfM1)} logistics / work-order / line aléas whose text reads like a quality problem, in 5 groups.",
  "Yes / no per group (no need to read them all).", "This workbook", "M1_QualityLike"),
 ("M4", "Medium", "ERMA aléas inside the disruption block", "19 of the 108 'no NC' quality aléas are on the MIP / 1101VB, where ERMA writes NCs. Their NC may exist but be excluded with ERMA.",
  "Put N in D_Block for those that have an ERMA NC.", "Skywise NCs on the same work order", "D_Block"),
 ("M2", "Medium", "Rework orders naming NCs we do not have", f"{len(dfM2)} orders, about 27 h, name 2026 NCs that are not in our exports.",
  "Plant / liability / writer of those NCs: out of scope, or missing from the export?", "Skywise NC lookup", "M2_NCsNotHeld"),
 ("M3", "Medium", "Exceptional NCs (not for the presentation)", "4 NCs where the order plans far more than the model, e.g. 34 h of cable reprinting on one NC line.",
  "Is the planned time right? Should such NCs be costed from their order?", "You", "M3_Exceptional"),
 ("L1", "Low", "MSN 700/701 work (IEVB)", "74 orders, 217 h in September, flight-test installation work.", "Confirm out of scope.", "Flight-test / IEV team", "Later"),
 ("L2", "Low", "Change 5 kept for later", "Two defect classes missing from the time table (about 0 EUR).", "Apply when convenient.", "-", "Later"),
]
hdr = ["ID", "Priority", "Check", "What we know", "Question to answer", "Where to look / who to ask", "Tab", "Status", "Finding", "Decision"]
dfC = pd.DataFrame(items, columns=hdr[:7]); dfC["Status"] = "To do"; dfC["Finding"] = ""; dfC["Decision"] = ""
r = para(wsC, 4, "Yellow cells are yours to fill (Status, Finding, Decision here; the yellow columns in each tab). Example of a filled line: Status 'Done' · Finding 'Orders are created by "
         "support without times when the rework is done by the test team' · Decision 'Use the model for orders with no time'.", ncols=10)
c0, c1 = table(wsC, 6, dfC, [5, 8, 34, 60, 40, 30, 18, 11, 40, 30], inputs=["Status", "Finding", "Decision"], status=["Status"], filt=True)
for i in range(c0, c1 + 1):
    p = wsC.cell(row=i, column=2).value
    wsC.cell(row=i, column=2).fill = {"High": fill_hi, "Medium": fill_med, "Low": fill_low}[p]
    tab = wsC.cell(row=i, column=7)
    tab.hyperlink = f"#{tab.value}!A1"; tab.font = Font(name=F, size=10, color="0563C1", underline="single")
total_row(wsC, c1 + 2, "Progress", [(3, f'=COUNTIF(H{c0}:H{c1},"Done")&" done, "&COUNTIF(H{c0}:H{c1},"In progress")&" in progress, "&COUNTIF(H{c0}:H{c1},"To do")&" to do"', None)], 10)

# ================================================================== Read me
ws0 = wb["Sheet"]; ws0.title = "ReadMe"; ws0.sheet_view.showGridLines = False
ws0["A1"] = "CNQ investigation workbook"; ws0["A1"].font = f_title
ws0["A2"] = "FTO2 / Airbus Electric Center, quality. Built 19 October 2026."; ws0["A2"].font = f_sub
ws0.column_dimensions["A"].width = 150
lines = [
 ("Two calculations", True),
 ("1. NC cost of non-quality: one cost per NC (level 3 of the CNQ model). The operator's aléa that leads to an NC is part of it (the 30-min alert step).", False),
 ("2. Disruption block, shown beside it and never added inside it: quality aléas that did not lead to an NC (option B), plus other-kind aléas you decide count as quality (M1).", False),
 ("", False),
 ("How to use", True),
 ("Start with the Checklist tab: one line per check, priority, question, where to look, and a link to its tab. Each tab lists the orders, NCs or aléas to look at.", False),
 ("Yellow cells (blue text) are yours to fill: Y / N answers, findings, figures you find. Totals and the Figures tab recalculate from them.", False),
 ("Example: in D_Block, an aléa whose NC was written by ERMA gets N in 'Keep in block?'; the block total in Figures drops by its hours.", False),
 ("Example: in H3, typing 3 in 'Repair hours at the customer' replaces the model's rework for that return by 3 h x the rate in Figures!C4.", False),
 ("", False),
 ("Sources", True),
 ("Work orders: SAP operations confirmed 1 Sep - 1 Oct 2026 (order type S03 = rework orders). Planned time = standard value in IM (1/100 h); confirmed times equal planned for 98% of operations.", False),
 ("Disruptions: MES export, 2,500 rows, 5 Jan - 30 Sep 2026. Link to NCs: same work order, NC created between 2 h before and 72 h after the aléa.", False),
 ("NCs: 2026 dashboard export + missing IDs (to 28 Sep 2026); 2025 NC file (plant FTO2). Rework order of an NC: AW_/WO_/RS_ number in linked_document different from its own work order.", False),
 ("CNQ model: cnq_handoever, version cnq-2026.10.19. Scripts that rebuild every list: 07_scripts/wo_disruption (build_checklist.py builds this file).", False),
]
for i, (s, b) in enumerate(lines, 4):
    c = ws0.cell(row=i, column=1, value=s); c.font = f_bold if b else f_base; c.alignment = Alignment(wrap_text=True)

order = ["ReadMe", "Checklist", "Figures", "H1_ReworkNoNC", "H2_QualityTime", "H3_CustomerReturns", "H4_CableCutting", "H5_NCsNoOrder", "H6_OrdersNoTime",
         "H7_DisruptExport", "D_Block", "M1_QualityLike", "M2_NCsNotHeld", "M3_Exceptional", "Later"]
wb._sheets = [wb[n] for n in order]
for ws in wb.worksheets:
    ws.sheet_properties.tabColor = {"Checklist": "1F3864", "Figures": "1F3864"}.get(ws.title, "C00000" if ws.title.startswith("H") else "BF8F00" if ws.title[0] in "MD" else "808080")
wb.active = 1
wb.save(OUT)
print("saved", OUT, [ (ws.title, ws.max_row) for ws in wb.worksheets ])
